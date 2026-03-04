"""
Word (.docx) audit report builder for Movern.

Usage::

    from movern.ui.report.docx_builder import build_report
    report_bytes = build_report(results, metadata)
    # pass report_bytes to st.download_button
"""

import io
from datetime import date
from typing import Any, Dict, List, Optional

from movern.ui.report.standards_map import STANDARDS_MAP


def build_report(
    results: List[Dict[str, Any]],
    metadata: Dict[str, Any],
    evaluator_names: Optional[List[str]] = None,
) -> bytes:
    """Build a Word audit report and return it as bytes.

    Parameters
    ----------
    results : list of dict
        Each dict has keys: ``metric``, ``value``, ``group`` (optional),
        ``evaluator``.
    metadata : dict
        Model metadata from ``load_demo_model()`` or user config.
    evaluator_names : list of str, optional
        Names of evaluators that were run.

    Returns
    -------
    bytes
        Raw .docx file bytes suitable for ``st.download_button``.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
    except ImportError:
        raise ImportError(
            "python-docx is required for report generation. "
            "Install it with: pip install python-docx"
        )

    doc = Document()

    # ---- Styles ----
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def heading(text, level=1, color=None):
        h = doc.add_heading(text, level=level)
        if color:
            for run in h.runs:
                run.font.color.rgb = RGBColor(*color)
        return h

    def add_table_row(table, cells):
        row = table.add_row()
        for i, val in enumerate(cells):
            row.cells[i].text = str(val)
        return row

    def shade_row(row, hex_color="D9E1F2"):
        for cell in row.cells:
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), hex_color)
            shd.set(qn("w:val"), "clear")
            tcPr.append(shd)

    # =====================================================================
    # COVER PAGE
    # =====================================================================
    doc.add_paragraph()
    cover_title = doc.add_paragraph()
    cover_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cover_title.add_run("Movern AI Governance Assessment Report")
    run.bold = True
    run.font.size = Pt(22)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(f"Model: {metadata.get('name', 'Unknown')}\n").bold = True
    p.add_run(f"Assessment Date: {date.today().isoformat()}\n")
    if evaluator_names:
        p.add_run(f"Evaluators: {', '.join(evaluator_names)}\n")
    p.add_run(f"EU AI Act Risk Tier: {metadata.get('eu_ai_act_risk', 'Not assessed')}\n")

    doc.add_page_break()

    # =====================================================================
    # 1. MODEL OVERVIEW
    # =====================================================================
    heading("1. Model Overview")
    info = doc.add_paragraph()
    info.add_run("Description: ").bold = True
    info.add_run(metadata.get("description", ""))
    doc.add_paragraph()
    info2 = doc.add_paragraph()
    info2.add_run("Features: ").bold = True
    info2.add_run(", ".join(metadata.get("features", [])))
    info3 = doc.add_paragraph()
    info3.add_run("Sensitive Features: ").bold = True
    info3.add_run(", ".join(metadata.get("sensitive_features", [])) or "None specified")
    info4 = doc.add_paragraph()
    info4.add_run("EU AI Act Risk Tier: ").bold = True
    info4.add_run(metadata.get("eu_ai_act_risk", "Not assessed"))
    info5 = doc.add_paragraph()
    info5.add_run("Applicable EU AI Act Articles: ").bold = True
    info5.add_run(", ".join(metadata.get("eu_ai_act_articles", [])))

    # =====================================================================
    # 2. ASSESSMENT RESULTS
    # =====================================================================
    heading("2. Assessment Results")

    # Group results by evaluator
    by_evaluator: Dict[str, List] = {}
    for r in results:
        ev = r.get("evaluator", "General")
        by_evaluator.setdefault(ev, []).append(r)

    for ev_name, ev_results in by_evaluator.items():
        heading(ev_name, level=2)

        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Metric"
        hdr[1].text = "Value"
        hdr[2].text = "Group"
        shade_row(table.rows[0], "2E75B6")
        for cell in table.rows[0].cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

        for r in ev_results:
            val = r.get("value")
            if isinstance(val, float):
                val = f"{val:.4f}"
            add_table_row(table, [
                r.get("metric", ""),
                val if val is not None else "—",
                r.get("group", "Overall"),
            ])
        doc.add_paragraph()

    # =====================================================================
    # 3. COMPLIANCE MAPPING
    # =====================================================================
    heading("3. Compliance Mapping")
    doc.add_paragraph(
        "The table below maps each assessed metric to relevant regulatory "
        "and standards requirements."
    )

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, col in enumerate(["Metric", "Value", "EU AI Act", "NIST AI RMF", "ISO 42001"]):
        hdr[i].text = col
    shade_row(table.rows[0], "2E75B6")
    for cell in table.rows[0].cells:
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    # De-duplicate by metric key
    seen = set()
    for r in results:
        key = r.get("metric", "")
        if key in seen:
            continue
        seen.add(key)
        mapping = STANDARDS_MAP.get(key, {})
        val = r.get("value")
        if isinstance(val, float):
            val = f"{val:.4f}"
        add_table_row(table, [
            mapping.get("display", key),
            val if val is not None else "—",
            mapping.get("eu_ai_act", "—"),
            mapping.get("nist_ai_rmf", "—"),
            mapping.get("iso_42001", "—"),
        ])

    doc.add_paragraph()

    # =====================================================================
    # 4. APPENDIX
    # =====================================================================
    heading("4. Appendix — Attribution")
    doc.add_paragraph(
        "This report was generated by Movern (https://github.com/movern-ai/movern), "
        "a community fork of Lens by Credo AI "
        "(https://github.com/credo-ai/credoai_lens). "
        "The assessment engine and evaluators originate from that project. "
        "Original authors: Ian Eisenberg and the Credo AI team. "
        "License: Apache 2.0."
    )

    # ---- Serialize to bytes ----
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
